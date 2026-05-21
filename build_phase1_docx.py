from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RED = RGBColor(0xc8, 0x32, 0x1a)
GOLD = RGBColor(0xc9, 0xa8, 0x4c)
BLACK = RGBColor(0x0d, 0x0d, 0x0d)
GRAY = RGBColor(0x6b, 0x6b, 0x6b)
WHITE = RGBColor(0xff, 0xff, 0xff)
CREAM_HEX = "F5F2EE"
RED_HEX = "C8321A"
GOLD_HEX = "C9A84C"
BLACK_HEX = "0D0D0D"

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.font.color.rgb = BLACK

for s in doc.sections:
    s.top_margin = Cm(2)
    s.bottom_margin = Cm(2)
    s.left_margin = Cm(2)
    s.right_margin = Cm(2)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def add_title(text, size=28, color=RED, bold=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return p


def add_subtitle(text, color=GRAY, size=11, italic=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.color.rgb = color
    return p


def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = BLACK
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:color'), RED_HEX)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_h3(text, color=RED):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = color
    return p


def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Cm(0.6)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.bold = True
        r.font.color.rgb = BLACK
        r2 = p.add_run(" " + text)
        r2.font.color.rgb = BLACK
    else:
        r = p.add_run(text)
        r.font.color.rgb = BLACK
    return p


def add_table(headers, rows, header_fill=BLACK_HEX, even_fill=CREAM_HEX, header_color=WHITE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        shade(hdr[i], header_fill)
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.color.rgb = header_color
                run.font.size = Pt(10.5)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        fill = even_fill if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row):
            cells[j].text = str(val)
            shade(cells[j], fill)
            for para in cells[j].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = BLACK
    return table


def add_pricing_table(headers, rows):
    """Pricing table with red header and gold-highlighted price cells."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        shade(hdr[i], RED_HEX)
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.color.rgb = WHITE
                run.font.size = Pt(10.5)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        fill = CREAM_HEX if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row):
            cells[j].text = str(val)
            cell_fill = "FFF6D9" if j == len(headers) - 1 else fill
            shade(cells[j], cell_fill)
            for para in cells[j].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    if j == len(headers) - 1:
                        run.font.bold = True
                        run.font.color.rgb = RED
                    else:
                        run.font.color.rgb = BLACK
    return table


add_title("Phase 1 — Scope & Pricing", size=30)
add_subtitle("Info For You Auto  ·  Prepared by Quad 4 Consulting", color=GRAY, size=11, italic=True)
add_subtitle("Phase 1 delivers brand assets and the public website with third-party rating integration. No internal certification system, no database, no admin panel — those are scoped into later phases.", color=GRAY, size=10, italic=True)

add_h2("1. Branding Deliverables")

add_h3("Logo Design")
add_bullet("3 design concepts presented for review")
add_bullet("2 rounds of revisions on the chosen concept")
add_bullet("Final files in vector (.svg, .ai)")
add_bullet("PNG with transparent background (multiple sizes)")
add_bullet("Favicon for the website (.ico + .png)")

add_h3("Visiting Card Design")
add_bullet("Front and back layouts")
add_bullet("Print-ready PDF (300 DPI, CMYK, with bleed)")
add_bullet("High-resolution JPG for digital sharing")
add_bullet("Design matched to the approved logo and brand palette")

add_h2("2. Website Build")

add_h3("Home Page + Podcast Page")
add_bullet("Branded home page: hero, recent episodes, Car Finder explainer, dealer preview")
add_bullet("Podcast archive listing every episode")
add_bullet("Inline YouTube video player on episode cards (click-to-play)")

add_h3("Car Finder (7-Step Tool)")
add_bullet("Guided 7-step flow: Budget → Body Style → Fuel → Seats → Use → Priorities → ZIP")
add_bullet("Top-5 vehicle results with plain-English reasons for each pick")
add_bullet("Nearby dealer recommendations ranked by proximity + third-party ratings")
add_bullet("One-click 'Get directions' opens Google Maps")

add_h3("Dealer Apply Form")
add_bullet("Multi-step application form on /dealers/apply")
add_bullet("Live Google Places autocomplete for dealership lookup")
add_bullet("Auto-fill of address, phone, website from Google")
add_bullet("Optional Yelp business ID for second rating source")
add_bullet("Submission emails AutoInfo4U with full application + fetched ratings")
add_bullet("AutoInfo4U approves via email reply; dealer is manually added to the site")

add_h3("Third-Party Ratings on Dealer Cards")
add_bullet("Google Reviews — star rating + total review count, pulled via Google Places API")
add_bullet("Yelp — star rating + review count, pulled via Yelp Fusion API")
add_bullet("BBB rating — manually entered by AutoInfo4U during approval (A+, A, B+, etc.)")
add_bullet("Aggregate trust score computed from above signals")
add_bullet("All ratings displayed inline on dealer cards as independent social proof")

add_h3("Packages Page")
add_bullet("Public pricing page showing dealer listing fees and podcast sponsorship tiers")
add_bullet("Each pricing card links directly to the dealer application or sponsor inquiry form")

add_h3("Hosting, Domain, Source Code")
add_bullet("Deployed to Vercel with automatic HTTPS and global CDN")
add_bullet("Custom domain configuration (AutoInfo4U owns the domain)")
add_bullet("Full source code transferred to AutoInfo4U's GitHub repository")

add_h2("3. Pricing Structure  —  EDITABLE")

add_subtitle("Replace every [ENTER PRICE] cell with your actual pricing before sending to AutoInfo4U. Cells highlighted in gold are the editable price fields.", color=GRAY, size=10, italic=True)

add_h3("3.1  Phase 1 Build — One-Time Fees", color=RED)
add_pricing_table(
    ["#", "Service line", "Price (USD)"],
    [
        ("1", "Logo design (3 concepts + 2 revisions + all file formats)", "[ENTER PRICE]"),
        ("2", "Visiting card design (front + back, print-ready)", "[ENTER PRICE]"),
        ("3", "Strategy & brand foundation (discovery + brand tokens)", "[ENTER PRICE]"),
        ("4", "Website design & development (all pages, responsive)", "[ENTER PRICE]"),
        ("5", "Car Finder tool (7-step flow + scoring algorithm)", "[ENTER PRICE]"),
        ("6", "Dealer apply form + Google Places + Yelp integration", "[ENTER PRICE]"),
        ("7", "Third-party ratings on dealer cards (Google + Yelp + BBB)", "[ENTER PRICE]"),
        ("8", "Podcast presentation (cards + YouTube player)", "[ENTER PRICE]"),
        ("9", "Hosting setup, custom domain wiring, GitHub transfer", "[ENTER PRICE]"),
        ("10", "Documentation (README, brand guide, scaffolding commands)", "[ENTER PRICE]"),
        ("11", "60-min recorded handoff + 2 weeks of email support", "[ENTER PRICE]"),
        ("", "PHASE 1 BUILD TOTAL", "[ENTER TOTAL]"),
    ],
)

add_h3("3.2  Dealer Directory — Recurring Listing Fees", color=RED)
add_subtitle("No internal certification tiers in Phase 1. Dealer trust comes from real third-party ratings (Google + Yelp + BBB). Set a pricing model that fits your business:", color=GRAY, size=10, italic=True)

add_pricing_table(
    ["Option", "Description", "Price (USD)"],
    [
        ("A — Flat listing fee", "Every approved dealer pays the same monthly fee to be listed in the directory and matched in Car Finder.", "[ENTER PRICE / MONTH]"),
        ("B — Featured upgrade", "Base listing fee + optional 'Featured' placement on the home page and top of search results.", "Base: [ENTER PRICE / MONTH]\nFeatured upgrade: [ENTER PRICE / MONTH]"),
        ("C — Per-lead pricing", "Free listing, dealer pays per qualified lead delivered (requires lead capture — Phase 2 build).", "[ENTER PRICE / LEAD]"),
        ("D — Annual prepay", "Discounted annual rate paid upfront instead of monthly.", "[ENTER PRICE / YEAR]"),
    ],
)

add_h3("3.3  Podcast Sponsorship — Per-Engagement Fees", color=RED)
add_pricing_table(
    ["Tier", "What's included", "Price (USD)"],
    [
        ("Episode Spot", "60-second mid-roll read in one episode + show notes link + one social mention", "[ENTER PRICE / EPISODE]"),
        ("Series Sponsor", "Mid-roll in 4 consecutive episodes + pre-roll mention + logo on episode artwork", "[ENTER PRICE / 4-EPISODE ARC]"),
        ("Season Sponsor", "Title sponsor for a 12-episode season + pre- and mid-roll in every episode + featured placement on the site", "[ENTER PRICE / QUARTER]"),
    ],
)

add_h3("3.4  Optional Add-Ons", color=RED)
add_pricing_table(
    ["Add-on", "Description", "Price (USD)"],
    [
        ("Analytics setup (PostHog)", "Pageview tracking + Car Finder funnel + dealer click events + 3 pre-built dashboards", "[ENTER PRICE]"),
        ("Additional rating source", "Add a fourth review source (e.g., Facebook Reviews, DealerRater scraping)", "[ENTER PRICE]"),
        ("Extra logo concept rounds", "Each additional round of revision beyond the included 2 rounds", "[ENTER PRICE / ROUND]"),
        ("Letterhead + email signature design", "Print + digital stationery to match the logo", "[ENTER PRICE]"),
        ("Branded social media templates", "Instagram + Facebook + LinkedIn post templates in the brand style", "[ENTER PRICE]"),
    ],
)

add_h3("3.5  Payment Terms", color=RED)
add_pricing_table(
    ["Milestone", "Schedule"],
    [
        ("On signing", "[ENTER % OR FIXED AMOUNT] of Phase 1 build total"),
        ("On design approval (logo + cards + website mockup)", "[ENTER % OR FIXED AMOUNT]"),
        ("On final delivery and handoff", "Balance"),
        ("Dealer directory + sponsorship", "Billed monthly / per-engagement per AutoInfo4U's chosen pricing model"),
    ],
)

add_h2("4. What Happens After Phase 1 Goes Live")

add_h3("For each new dealer that applies")
add_bullet("Dealer fills the apply form on /dealers/apply", bold_prefix="1.")
add_bullet("Form auto-fills info via Google Places and fetches their live Google + Yelp ratings", bold_prefix="2.")
add_bullet("AutoInfo4U receives the application via email, including all third-party ratings", bold_prefix="3.")
add_bullet("AutoInfo4U vets BBB rating and any other manual checks", bold_prefix="4.")
add_bullet("AutoInfo4U replies to approve / request more info / reject", bold_prefix="5.")
add_bullet("Approved dealers are added to the site (manual code edit in Phase 1)", bold_prefix="6.")
add_bullet("Dealer card goes live showing real Google + Yelp + BBB ratings as social proof", bold_prefix="7.")

add_h3("For each new podcast episode")
add_bullet("Upload audio to Spotify for Podcasters (free) — distributes to Apple Podcasts, Spotify, etc.")
add_bullet("Upload video to YouTube (free)")
add_bullet("Add the YouTube video ID + audio URL to the site's episode data file")
add_bullet("Push to GitHub — episode appears on the site in ~60 seconds")

add_h2("5. Phase 1 Boundaries — For Clarity")

p = doc.add_paragraph()
r = p.add_run("These capabilities are intentionally NOT in Phase 1 — they are scoped into Phase 2 or Phase 3:")
r.font.size = Pt(10.5)
r.font.color.rgb = BLACK

excluded = [
    "Database for persistent storage of episodes, dealers, leads, applications",
    "Lead capture from Car Finder (storing shopper preferences and contact info)",
    "Email/SMS lead routing to matched dealers",
    "Post-engagement shopper surveys (NPS, response time, quote accuracy)",
    "Stripe-based payment collection for dealer subscriptions",
    "Login / authentication for dealers or admins",
    "Admin panel for managing dealers and reviewing applications",
    "Dealer self-service dashboard",
    "Automated daily refresh of Google/Yelp ratings",
    "Per-episode and per-dealer SEO pages",
    "Newsletter automation",
]
for item in excluded:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run(item)
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Quad 4 Consulting  ·  info@quad4consulting.com")
r.font.size = Pt(9)
r.font.color.rgb = GRAY
r.font.italic = True

out = "Phase1_Scope_AutoInfo4U.docx"
doc.save(out)
print("OK:", out)
